#!/usr/bin/env python3
"""
Grid Checker for the Agency Grid Processor (Special Motor Matrix).

Validates a converted grid CSV in up to three layers:

  1. OUTPUT checks   (always)            - the CSV on its own: schema, leftover '-', blanks,
                                           LL<UL, volume mapping, % scaling, span logic,
                                           product signatures, RTO consistency, duplicates.
  2. SOURCE recon    (--source --sheet)   - rebuilds the expected grid from the source Excel
                                           using ONLY the golden rules file (not the portal
                                           config) and compares cell by cell, with source
                                           cell references for every mismatch.
  3. CONFIG lint     (--config)           - checks a portal config JSON against the sheet and
                                           the rules, to catch the mistake before conversion.

Usage
  python grid_checker.py --csv final.csv --rules rules/special_comp_rules.json \
      [--source grid.xlsx --sheet "Special Matrix-Comp.-1-14"] [--config portal_cfg.json] \
      [--version-id agency_spl_sep26_Grid_3] [--report report.xlsx]

Exit code: 0 = no ERROR findings, 1 = at least one ERROR, 2 = checker could not run.
"""
import argparse, json, re, sys, os
from collections import defaultdict, Counter
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

ERROR, WARN, INFO = 'ERROR', 'WARN', 'INFO'
SENT_LL, SENT_UL = -999999999, 999999999
PLACEHOLDERS = {'-', '--', 'nan', 'NaN', 'None', 'null', 'NULL', '#N/A', 'N/A', '#REF!', '#VALUE!', '#DIV/0!', '#NAME?'}


def nkey(s):
    """Normalise a label for matching: upper-case, no whitespace."""
    return re.sub(r'\s+', '', str(s or '')).upper()


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def fmt_num(x):
    if x is None:
        return ''
    x = round(float(x), 6)
    return str(int(x)) if x == int(x) else repr(x)


def prep(df):
    """Fast per-frame normalisation using unique-value maps.
    Returns S (stripped strings), N (floats, NaN if not numeric), K (canonical: numbers as fmt_num, else S)."""
    S, N, K = {}, {}, {}
    for c in df.columns:
        col = df[c].astype(str)
        u = pd.unique(col)
        st = {x: x.strip() for x in u}
        nm = {x: to_num(st[x]) for x in u}
        S[c] = col.map(st)
        N[c] = col.map(lambda x, nm=nm: nm[x]).astype(float)
        K[c] = col.map({x: (fmt_num(nm[x]) if nm[x] is not None else st[x]) for x in u})
    return pd.DataFrame(S, index=df.index), pd.DataFrame(N, index=df.index), pd.DataFrame(K, index=df.index)


class Findings:
    def __init__(self, overrides=None):
        self.items = []      # dicts: layer, check_id, check, severity, count, column, expected, actual, examples, fix
        self.ov = {k: v for k, v in (overrides or {}).items() if isinstance(v, dict)}
        self.mismatch = []; self.meta = {}

    def add(self, layer, check_id, check, severity, count=1, column='', expected='', actual='', examples='', fix=''):
        o = self.ov.get(check_id, {})
        if o.get('enabled') is False:
            return
        if o.get('severity') in (ERROR, WARN, INFO):
            severity = o['severity']
        self.items.append(dict(layer=layer, check_id=check_id, check=check, severity=severity, count=int(count),
                               column=column, expected=str(expected)[:300], actual=str(actual)[:300],
                               examples=str(examples)[:500], fix=fix))

    def has_errors(self):
        return any(i['severity'] == ERROR for i in self.items)


# ════════════════════════════════════════════════════════════════════════════
# STD-GRID VOLUME RULES  (ordered, first match wins)
#   {"agent_group": "Agency" | "*", "biz_mix_match": any|equals|starts_with|contains|in_list|regex,
#    "biz_mix_value": "PVT CAR", "target": "Totalgwp Pvt Car", "enabled": true, "note": ""}
# ════════════════════════════════════════════════════════════════════════════

MATCH_TYPES = ('any', 'equals', 'starts_with', 'contains', 'in_list', 'regex')


def norm_ag(s):
    return re.sub(r'[^A-Z0-9]', '', str(s or '').upper())


def bm_match(kind, val, bm):
    b = str(bm or '').strip().upper(); v = str(val or '').strip().upper()
    if kind in ('any', '', None): return True
    if kind == 'equals': return b == v
    if kind == 'starts_with': return b.startswith(v)
    if kind == 'contains': return v in b
    if kind == 'in_list': return b in {x.strip().upper() for x in v.split(',') if x.strip()}
    if kind == 'regex': return re.search(str(val), str(bm or ''), re.I) is not None
    return False


def std_rules(rules):
    r = rules.get('std_grid_volume_rules')
    if r is None:   # legacy rules files: one pair per IMD type
        r = [{'agent_group': v.get('agent_group'), 'biz_mix_match': 'any', 'biz_mix_value': '', 'target': v['std_grid_volume']}
             for k, v in rules.get('imd_type_map', {}).items() if isinstance(v, dict) and v.get('std_grid_volume')]
    return [x for x in r if x.get('enabled', True) and x.get('target')]


def resolve_std(rl, agent_group, biz_mix):
    a = norm_ag(agent_group)
    for i, x in enumerate(rl):
        g = x.get('agent_group', '*')
        if g not in ('*', '', None) and norm_ag(g) != a:
            continue
        if bm_match(x.get('biz_mix_match', 'any'), x.get('biz_mix_value', ''), biz_mix):
            return x['target'], i
    return None, None


def describe_rule(x):
    m = x.get('biz_mix_match', 'any')
    cond = 'any Biz Mix' if m in ('any', '', None) else f"Biz Mix {m.replace('_', ' ')} '{x.get('biz_mix_value', '')}'"
    return f"{x.get('agent_group', '*')} + {cond} → {x['target']}"


def user_map(d):
    return {k: v for k, v in (d or {}).items() if not str(k).startswith('_')}


# ════════════════════════════════════════════════════════════════════════════
# SOURCE SHEET READING (independent of portal config — header-driven)
# ════════════════════════════════════════════════════════════════════════════

class SheetLayout:
    pass


def read_sheet(xlsx, sheet, rules):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet '{sheet}' not found. Available: {wb.sheetnames}")
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    lay = rules['sheet_layout']
    L = SheetLayout()
    L.rows = rows
    anchor = nkey(lay['anchor_label'])
    L.h3 = next((i for i, r in enumerate(rows[:30]) if any(nkey(v) == anchor for v in r)), None)
    if L.h3 is None:
        raise SystemExit(f"Could not find header row containing '{lay['anchor_label']}' in first 30 rows of '{sheet}'.")
    L.h2, L.h1, L.data_start = L.h3 - 1, L.h3 - 2, L.h3 + 1          # 0-based
    hdr3 = rows[L.h3]; hdr2 = rows[L.h2]; hdr1 = rows[L.h1]
    width = max(len(r) for r in rows)
    get = lambda r, c: (r[c] if c < len(r) else None)
    L.meta = {}
    for key, label in lay['meta_labels'].items():
        L.meta[key] = next((c for c in range(width) if nkey(get(hdr3, c)) == nkey(label)), None)
    L.pairs = {}
    for key, glabel in lay['lower_upper_groups'].items():
        c = next((c for c in range(width) if nkey(get(hdr2, c)) == nkey(glabel)), None)
        if c is not None and nkey(get(hdr3, c)) == 'LOWER' and nkey(get(hdr3, c + 1)) == 'UPPER':
            L.pairs[key] = (c, c + 1)
        else:
            L.pairs[key] = None
    L.rate_cols = []
    for c in range(width):
        basis = {nkey(k): v for k, v in rules['rate_cell']['outgo_by_header_label'].items()}.get(nkey(get(hdr3, c)))
        if basis:
            L.rate_cols.append(dict(col=c, parent=str(get(hdr1, c) or '').strip(), sub=str(get(hdr2, c) or '').strip(),
                                    label=str(get(hdr3, c)).strip(), outgo=basis))
    return L


def read_rto_master(xlsx, rules):
    r = rules['rto']
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if r['sheet'] not in wb.sheetnames:
        wb.close(); return None
    rows = list(wb[r['sheet']].iter_rows(values_only=True)); wb.close()
    hdr = [str(v or '').strip().upper() for v in rows[r['header_row'] - 1]]
    cc = hdr.index(r['code_col'].upper()) if r['code_col'].upper() in hdr else None
    kc = next((i for i, h in enumerate(hdr) if h.startswith(r['cluster_col_prefix'].upper())), None)
    if cc is None or kc is None:
        return None
    idx = defaultdict(set)
    for row in rows[r['header_row']:]:
        if cc < len(row) and kc < len(row) and row[cc] and row[kc]:
            idx[str(row[kc]).strip().upper()].add(str(row[cc]).strip())
    return {k: ','.join(sorted(v)) for k, v in idx.items()}


# ════════════════════════════════════════════════════════════════════════════
# LAYER 2 — BUILD EXPECTED GRID FROM SOURCE + RULES
# ════════════════════════════════════════════════════════════════════════════

def build_expected(L, sheet, rules, rto_master, F, version_id=None):
    rc = rules['rate_cell']; pb = rules['percent_band']
    irda_vals = {nkey(v) for v in rc['irda_trigger_values']}
    cmap = {(nkey(m['parent']), nkey(m['sub'])): m for m in rules['column_map']}
    vmap = {nkey(k): v for k, v in user_map(rules['volume_consideration_map']).items()}
    bmap = {nkey(k): v for k, v in user_map(rules['bizmix_consideration_map']).items()}
    imap = {nkey(k): v for k, v in user_map(rules['imd_type_map']).items()}
    SR = std_rules(rules); unk_std = defaultdict(list)
    skip_vals = {nkey(v) for v in rules['sheet_layout'].get('skip_rows_where_vol_lower_is', [])}
    cols = rules['expected_columns']

    # R02: header / column map coverage
    active = []
    for rc_ in L.rate_cols:
        m = cmap.get((nkey(rc_['parent']), nkey(rc_['sub'])))
        if not m:
            F.add('SOURCE', 'R02', 'Rate column in sheet has no rule (new / renamed LOB column)', ERROR, 1,
                  column=f"{get_column_letter(rc_['col']+1)}: {rc_['parent']} | {rc_['sub']}",
                  fix='Add this header to column_map in the rules file AND to Step 3 of the portal config.')
        else:
            active.append((rc_, m))
    seen = {(nkey(r['parent']), nkey(r['sub'])) for r in L.rate_cols}
    for k, m in cmap.items():
        if k not in seen:
            F.add('SOURCE', 'R02b', 'Rule column not present in this sheet', INFO, 1, column=f"{m['parent']} | {m['sub']}")
    for key in ('imd_code', 'imd_type', 'uw_cluster', 'vol_remark'):
        if L.meta.get(key) is None:
            F.add('SOURCE', 'R00', f'Meta column not found in sheet: {key}', ERROR, 1,
                  fix='Header label changed? Update sheet_layout.meta_labels in rules.')

    def g(row, key):
        c = L.meta.get(key)
        v = row[c] if c is not None and c < len(row) else None
        return '' if v is None else (v.strip() if isinstance(v, str) else v)

    def gp(row, key):
        p = L.pairs.get(key)
        if not p:
            return None, None
        return (row[p[0]] if p[0] < len(row) else None), (row[p[1]] if p[1] < len(row) else None)

    out, refs, vol_unmapped = [], [], []
    dup_src = defaultdict(list)
    unk_cells = defaultdict(list); unk_vol = defaultdict(list); unk_imd = defaultdict(list); unk_bm = defaultdict(list)
    for ri in range(L.data_start, len(L.rows)):
        row = L.rows[ri]; xrow = ri + 1
        imd_code, imd_name, cluster = str(g(row, 'imd_code')), str(g(row, 'imd_name')), str(g(row, 'uw_cluster'))
        vll, vul = gp(row, 'vol')
        if not any([imd_code, imd_name, cluster, '' if vll is None else str(vll)]):
            continue
        if nkey(vll) in skip_vals:
            continue
        imd_type = str(g(row, 'imd_type'))
        dup_src[(imd_code, imd_type, str(vll), str(vul), str(g(row, 'vol_remark')), cluster)].append(xrow)
        base = {c: rules['defaults'].get(c, '') for c in cols}
        if version_id:
            base['Version Id*'] = version_id
        base['Parent Agent Code*'] = imd_code or 'ANY'
        rel = str(g(row, 'rel_code'))
        if rel and rel != '-':
            base['Primary Agent Code*'] = rel
        im = imap.get(nkey(imd_type))
        if im:
            base['Agent Group Code*'] = im['agent_group']
        else:
            unk_imd[imd_type].append(xrow)
        base['Rto Cluster*'] = cluster or 'ANY'
        if rto_master is not None and cluster:
            base['Rto Code*'] = rto_master.get(cluster.upper(), '<<CLUSTER NOT IN RTO MASTER>>')

        # volume (STD-GRID is resolved per column below, because it depends on Biz Mix)
        vrem = nkey(g(row, 'vol_remark'))
        vpair, unmapped, is_std = None, False, False
        if to_num(vll) is not None or to_num(vul) is not None:
            if vrem in ('', '-'):
                unk_vol['(blank, but volume given)'].append(xrow); unmapped = True
            elif vrem == 'STD-GRID':
                is_std = True
            else:
                vpair = vmap.get(vrem)
                if vpair is None:
                    unk_vol[str(g(row, 'vol_remark'))].append(xrow); unmapped = True

        # discount band
        dl, du = gp(row, 'discount')
        if to_num(dl) is not None or to_num(du) is not None:
            p = rules['discount_pair']
            if to_num(dl) is not None: base[p + ' Ll*'] = fmt_num(to_num(dl) * pb['scale'])
            if to_num(du) is not None: base[p + ' Ul*'] = fmt_num(to_num(du) * pb['scale'] + pb['upper_add'])

        # biz-mix % band
        bl, bu = gp(row, 'bizmix')
        blabel = str(g(row, 'bizmix_label'))
        if to_num(bl) is not None or to_num(bu) is not None:
            nl = nkey(blabel)
            p = next((v for k, v in bmap.items() if nl.startswith(k) or k in nl), None)
            if p is None:
                unk_bm[blabel or '(blank label)'].append(xrow)
            else:
                if to_num(bl) is not None: base[p + ' Ll*'] = fmt_num(to_num(bl) * pb['scale'])
                if to_num(bu) is not None: base[p + ' Ul*'] = fmt_num(to_num(bu) * pb['scale'] + pb['upper_add'])

        for rc_, m in active:
            c = rc_['col']
            raw = row[c] if c < len(row) else None
            if raw is None or (isinstance(raw, str) and raw.strip() == ''):
                continue
            o = dict(base)
            o['Biz Mix*'] = m['biz_mix']
            o.update(m['fields'])
            cell_unmapped, p = unmapped, vpair
            if is_std:
                p, _ = resolve_std(SR, o.get('Agent Group Code*'), m['biz_mix'])
                if p is None:
                    unk_std[(o.get('Agent Group Code*'), m['biz_mix'])].append(xrow); cell_unmapped = True
            if p:
                if to_num(vll) is not None: o[p + ' Ll*'] = fmt_num(to_num(vll))
                if to_num(vul) is not None: o[p + ' Ul*'] = fmt_num(to_num(vul))
            if isinstance(raw, str) and nkey(raw) in {nkey(x) for x in rc['irda_trigger_values']}:
                o['Span Outgo*'] = rc['irda_outgo']; o['Span Prct*'] = fmt_num(rc['irda_prct'])
            elif to_num(raw) is not None:
                o['Span Outgo*'] = rc_['outgo']; o['Span Prct*'] = fmt_num(to_num(raw) * rc['scale'])
            else:
                unk_cells[str(raw)].append(f"{get_column_letter(c+1)}{xrow}")
                continue
            out.append(o)
            refs.append(f"'{sheet}'!{get_column_letter(c+1)}{xrow}")
            vol_unmapped.append(cell_unmapped)

    d = {k: v for k, v in dup_src.items() if len(v) > 1}
    if d:
        imds = Counter(k[0] for k in d)
        F.add('SOURCE', 'R10', 'Same IMD + cluster + volume slab appears more than once in the source sheet', WARN, len(d),
              actual='; '.join(f'{i} ({n} clusters)' for i, n in imds.most_common(5)),
              examples='; '.join(f"{k[0]}/{k[5]}: rows {v}" for k, v in list(d.items())[:5]),
              fix='Ask the grid owner which block is valid; if rates differ the output will contain conflicting rows (S14b).')
    for v, locs in unk_cells.items():
        F.add('SOURCE', 'R01', 'Unrecognised rate cell value in source (not a number, not an IRDA trigger)', ERROR, len(locs),
              actual=v, examples=', '.join(locs[:8]),
              fix='Either add the value to irda_trigger_values (rules + portal Step 2) or correct the source cell.')
    for v, locs in unk_vol.items():
        F.add('SOURCE', 'R03', 'Volume Consideration has no GWP-column rule', ERROR, len(locs), actual=v,
              examples='source rows ' + ', '.join(map(str, locs[:10])),
              fix='Add it to volume_consideration_map (rules) and vgRows (portal Step 4). Without it the portal silently writes to Total Gwp.')
    for (ag, bm), locs in unk_std.items():
        F.add('SOURCE', 'R03d', 'STD-GRID row: no volume rule matches this Agent Group + Biz Mix', ERROR, len(locs),
              actual=f'{ag} + {bm}', examples='source rows ' + ', '.join(map(str, sorted(set(locs))[:10])),
              fix='Add a STD-GRID volume rule (Checker Rules → STD-GRID rules).')
    for v, locs in unk_imd.items():
        F.add('SOURCE', 'R03b', 'IMD Type has no Agent Group rule', ERROR, len(locs), actual=v or '(blank)',
              examples='source rows ' + ', '.join(map(str, locs[:10])))
    for v, locs in unk_bm.items():
        F.add('SOURCE', 'R03c', 'Biz Mix consideration label has no Prct-Vol rule', ERROR, len(locs), actual=v,
              examples='source rows ' + ', '.join(map(str, locs[:10])))
    exp = pd.DataFrame(out, columns=cols).fillna('')
    return exp, refs, np.array(vol_unmapped, dtype=bool)


# ════════════════════════════════════════════════════════════════════════════
# MATCH EXPECTED <-> ACTUAL AND DIFF
# ════════════════════════════════════════════════════════════════════════════

def vol_signature(N, rules):
    """Value-only signature of all volume columns (which column they sit in is ignored)."""
    parts = []
    for p in rules['volume_pairs']:
        for suf, sent in ((' Ll*', SENT_LL), (' Ul*', SENT_UL)):
            if p + suf in N.columns:
                v = N[p + suf].to_numpy()
                parts.append(np.where(v == sent, np.nan, v))
    m = np.sort(np.vstack(parts).T, axis=1)          # NaN sorts last
    out = pd.Series([''] * len(m), index=N.index)
    for j in range(m.shape[1]):
        col = pd.Series(m[:, j], index=N.index)
        if col.notna().any():
            u = {x: fmt_num(x) for x in pd.unique(col.dropna())}
            out = out + '|' + col.map(u).fillna('')
    return out


def reconcile(exp, refs, vol_unmapped, act, rules, F, version_id):
    cols = [c for c in rules['expected_columns'] if c in act.columns]
    base_key = ['Parent Agent Code*', 'Agent Group Code*', 'Rto Cluster*', 'Biz Mix*']
    attr = [c for c in rules['product_attribute_columns'] if c in act.columns]
    eS, eN, eK = prep(exp[cols]); aS, aN, aK = prep(act[cols])
    def keys(K, N):
        k1 = K[base_key + attr].agg('\x1f'.join, axis=1) + '\x1f' + vol_signature(N, rules)
        k2 = K[base_key].agg('\x1f'.join, axis=1)
        return k1, k2
    e = pd.DataFrame({'_ei': np.arange(len(exp))}); a = pd.DataFrame({'_ai': np.arange(len(act))})
    e['_k1'], e['_k2'] = [x.to_numpy() for x in keys(eK, eN)]
    a['_k1'], a['_k2'] = [x.to_numpy() for x in keys(aK, aN)]
    pairs = []
    em, am = e, a
    for k in ('_k1', '_k2'):
        em = em.copy(); am = am.copy()
        em['_n'] = em.groupby(k).cumcount(); am['_n'] = am.groupby(k).cumcount()
        m = em[[k, '_n', '_ei']].merge(am[[k, '_n', '_ai']], on=[k, '_n'])
        pairs.append(m[['_ei', '_ai']])
        em = em[~em['_ei'].isin(m['_ei'])]; am = am[~am['_ai'].isin(m['_ai'])]
    P = pd.concat(pairs, ignore_index=True)
    ei, ai = P['_ei'].to_numpy(), P['_ai'].to_numpy()

    if len(exp) == len(act):
        F.add('SOURCE', 'R04i', 'Row count matches source', INFO, 0, expected=len(exp), actual=len(act))
    else:
        F.add('SOURCE', 'R04', 'Row count: expected (rebuilt from source) vs actual', ERROR, abs(len(exp) - len(act)),
              expected=len(exp), actual=len(act),
              fix='Rows missing/extra — see R05/R06. Usually a disabled column, wrong data start row, or rows deleted/duplicated manually.')
    if len(em):
        smp = em.head(8)
        F.add('SOURCE', 'R05', 'Expected rows missing from output', ERROR, len(em),
              examples='; '.join(f"{refs[i]} [{exp.at[i,'Parent Agent Code*']}/{exp.at[i,'Rto Cluster*']}/{exp.at[i,'Biz Mix*']}]" for i in smp['_ei']))
    if len(am):
        smp = am.head(8)
        F.add('SOURCE', 'R06', 'Output rows with no matching source cell (extra / changed identity)', ERROR, len(am),
              examples='; '.join(f"csv line {i+2} [{act.at[i,'Parent Agent Code*']}/{act.at[i,'Rto Cluster*']}/{act.at[i,'Biz Mix*']}]" for i in smp['_ai']))

    EK = eK.iloc[ei].reset_index(drop=True); AK = aK.iloc[ai].reset_index(drop=True)
    vu = vol_unmapped[ei]
    vol_cols = {p + s_ for p in rules['volume_pairs'] for s_ in (' Ll*', ' Ul*')}
    mismatch_rows = []
    for c in cols:
        if c == 'Version Id*' and not version_id:
            continue
        diff = (EK[c] != AK[c]).to_numpy()
        if c in vol_cols:
            diff = diff & ~vu
        if not diff.any():
            continue
        d = pd.DataFrame({'exp': EK[c][diff].to_numpy(), 'act': AK[c][diff].to_numpy(), 'ei': ei[diff], 'ai': ai[diff]})
        for (xv, av), grp in d.groupby(['exp', 'act'], sort=False):
            ex = ', '.join(f"{refs[i]}→csv line {j+2}" for i, j in zip(grp['ei'][:4], grp['ai'][:4]))
            F.add('SOURCE', 'R07', 'Cell value differs from source/rules', ERROR, len(grp), column=c,
                  expected=xv, actual=av, examples=ex, fix=_fix_hint(c, xv, av, rules))
        for i, j in zip(d['ei'][:3], d['ai'][:3]):
            mismatch_rows.append(dict(source_cell=refs[i], csv_line=int(j) + 2, column=c,
                                      expected=exp.at[i, c], actual=act.at[j, c]))
    # what did the output do with rows whose volume consideration has no rule?
    if vu.any():
        sub = act.iloc[ai[vu]]
        placed = Counter()
        for p in rules['volume_pairs']:
            v = pd.to_numeric(sub.get(p + ' Ll*', pd.Series(dtype=str)), errors='coerce')
            placed[p] += int(((v != SENT_LL) & v.notna()).sum())
        txt = ', '.join(f'{k}: {v} rows' for k, v in placed.items() if v) or 'nowhere (volume dropped)'
        F.add('SOURCE', 'R08', 'Rows with unruled Volume Consideration — where the output put the volume', WARN, int(vu.sum()),
              actual=txt, fix='Confirm the correct target column and add it to volume_consideration_map.')
    return mismatch_rows


def _fix_hint(col, xv, av, rules):
    if av in PLACEHOLDERS or av == '':
        return "Leftover '-'/blank from the source — the replace-with-default manual step was missed, or the column default is missing in portal Step 6."
    if col.startswith(tuple(rules['volume_pairs'])):
        return 'Volume mapped to the wrong column or wrong value — check vgRows / imdGwpRows (Step 4). Look for an LL column used as UL.'
    if col.startswith(tuple(rules['percent_pairs'])):
        try:
            if abs(float(av) * 100 + (0.001 if 'Ul*' in col else 0) - float(xv)) < 1e-6:
                return 'Fraction not converted to percent (×100, UL +0.001) — manual step missed.'
        except ValueError:
            pass
        return 'Wrong % band — check extraMetaCols mapping (Biz Mix consideration decides which Prct Vol pair).'
    if col == 'Span Prct*':
        return 'Rate differs — check ×100 transform, IRDA value, or column shift in Step 3.'
    if col == 'Rto Code*':
        return 'RTO codes differ — check RTO sheet/cluster column in Step 5.'
    return 'Check column mapping / extra fields for this Biz Mix in Step 3.'


# ════════════════════════════════════════════════════════════════════════════
# LAYER 1 — STANDALONE OUTPUT CHECKS
# ════════════════════════════════════════════════════════════════════════════

def check_output(act, rules, F, rto_master=None, version_id=None):
    L = 'OUTPUT'
    exp_cols = rules['expected_columns']
    missing = [c for c in exp_cols if c not in act.columns]
    extra = [c for c in act.columns if c not in exp_cols]
    if missing: F.add(L, 'S01', 'Expected columns missing', ERROR, len(missing), actual=', '.join(missing),
                      fix='Add them to Output Format (Step 6) and Column Defaults.')
    if extra: F.add(L, 'S01b', 'Unexpected columns present', ERROR, len(extra), actual=', '.join(extra))
    if not missing and not extra and list(act.columns) != exp_cols:
        F.add(L, 'S01c', 'Column order differs from template', WARN, 1, fix='Re-order via Output Format (Step 6).')
    s, NN, KK = prep(act)
    for c in act.columns:
        blank = (s[c] == '').sum()
        if blank:
            F.add(L, 'S02', 'Blank cells', ERROR, blank, column=c, examples=_lines(s.index[s[c] == '']),
                  fix='Add a Column Default for this column.')
        ph = s[c].isin(PLACEHOLDERS)
        if ph.any():
            F.add(L, 'S03', "Placeholder values left ('-', nan, #N/A …)", ERROR, int(ph.sum()), column=c,
                  actual=', '.join(sorted(set(s[c][ph]))), examples=_lines(s.index[ph]),
                  fix='Replace with the column default (manual step) — or set the default in Step 6.')
    num_cols = [c for c in act.columns if c.endswith(' Ll*') or c.endswith(' Ul*') or c == 'Span Prct*']
    N = NN[num_cols]
    for c in num_cols:
        bad = N[c].isna() & ~s[c].isin(PLACEHOLDERS) & (s[c] != '')
        if bad.any():
            F.add(L, 'S04', 'Non-numeric value in numeric column', ERROR, int(bad.sum()), column=c,
                  actual=', '.join(sorted(set(s[c][bad]))[:6]), examples=_lines(s.index[bad]))
        # float artefacts e.g. 28.999999999999996
        md = rules['rate_cell'].get('max_decimals', 4)
        u = {x: bool(re.search(r'\.\d{%d,}$' % (md + 1), x)) for x in pd.unique(s[c])}
        art = s[c].map(u)
        if art.any():
            F.add(L, 'S04b', 'Floating-point artefacts (too many decimals)', WARN, int(art.sum()), column=c,
                  actual=', '.join(sorted(set(s[c][art]))[:4]), examples=_lines(s.index[art]),
                  fix='Round to 2–4 decimals (portal ×100 transform creates these).')
    # LL / UL pairs
    bases = sorted({c[:-4] for c in act.columns if c.endswith(' Ll*') and c[:-4] + ' Ul*' in act.columns})
    for b in bases:
        ll, ul = N.get(b + ' Ll*'), N.get(b + ' Ul*')
        if ll is None or ul is None: continue
        bad = ll.notna() & ul.notna() & (ll >= ul)
        if bad.any():
            pairs = Counter(zip(s.loc[bad, b + ' Ll*'], s.loc[bad, b + ' Ul*']))
            F.add(L, 'S05', 'Lower limit >= Upper limit (slab can never match)', ERROR, int(bad.sum()), column=b,
                  actual='; '.join(f'LL={x} UL={y} ×{k}' for (x, y), k in pairs.most_common(4)), examples=_lines(s.index[bad]),
                  fix='Usually an LL column configured as the UL target (Step 4), or LL/UL swapped.')
    # volume: at most one volume pair set per row
    vp = [p for p in rules['volume_pairs'] if p + ' Ll*' in N.columns]
    setm = pd.concat([((N[p + ' Ll*'].notna() & (N[p + ' Ll*'] != SENT_LL)) |
                       (N[p + ' Ul*'].notna() & (N[p + ' Ul*'] != SENT_UL))).rename(p) for p in vp], axis=1)
    multi = setm.sum(axis=1) > 1
    if multi.any():
        combos = Counter(setm[multi].apply(lambda r: ' + '.join(r.index[r]), axis=1))
        F.add(L, 'S06', 'More than one volume (GWP) column populated in a row', ERROR, int(multi.sum()),
              actual='; '.join(f'{k} ×{v}' for k, v in combos.most_common(4)), examples=_lines(s.index[multi]))
    for p in vp:
        llbad = N[p + ' Ll*'] == SENT_UL
        if llbad.any():
            F.add(L, 'S07', 'Volume LL holds the UL sentinel 999999999', ERROR, int(llbad.sum()), column=p + ' Ll*',
                  examples=_lines(s.index[llbad]), fix='UL value was written into the LL column (imdGwpRows / vgRows ul_col points at an Ll* column).')
        ulbad = N[p + ' Ul*'] == SENT_LL
        if ulbad.any():
            F.add(L, 'S07', 'Volume UL holds the LL sentinel -999999999', ERROR, int(ulbad.sum()), column=p + ' Ul*', examples=_lines(s.index[ulbad]))
    # volume vs agent group (std-grid rows) and vs biz mix (product-specific volume columns)
    if 'Agent Group Code*' in s.columns:
        tab = pd.crosstab(s['Agent Group Code*'], setm.idxmax(axis=1).where(setm.any(axis=1), '(none)'))
        F.add(L, 'S06i', 'Volume column usage by Agent Group (review)', INFO, 0,
              actual='; '.join(f"{ag}: " + ', '.join(f'{c}={v}' for c, v in r.items() if v) for ag, r in tab.iterrows()))
    prod_vol = rules.get('product_volume_columns') or {'Totalgwp Tractor': ['TRACTOR'], 'Totalgwp Tractorold': ['TRACTOR'], 'Totalgwp Tractornew': ['TRACTOR'],
                'Totalgwp Pvt Car': ['PVT CAR'], 'Totalgwp Pvtcar1plus1': ['PVT CAR'], 'Two Wheeler 1 O 1': ['2W'],
                'Totalgwp 20 40t Comp': ['GCV'], 'Totalgwp 12 20t Comp': ['GCV'], 'Totalgwp 7500 12000': ['GCV'],
                'Totalgwp Gcvgreater40t': ['GCV'], 'Tractor Tp Total Gwp': ['TRACTOR'], 'Gcv Tpless2500 Total Gwp': ['GCV']}
    prod_vol = user_map(prod_vol)
    for p, toks in prod_vol.items():
        toks = [t.upper() for t in toks]
        if p not in setm.columns: continue
        on = setm[p]
        if not on.any(): continue
        bm = s.loc[on, 'Biz Mix*'].str.upper()
        odd = ~bm.apply(lambda x: any(t in x for t in toks))
        # product-specific volume on another LOB is legitimate only if the whole IMD is volume-tested on it;
        # flag when the IMD has NO row of the matching LOB at all
        if odd.any():
            imds = s.loc[on[on].index[odd.values], 'Parent Agent Code*']
            sus = []
            for imd in imds.unique():
                lobs = s.loc[s['Parent Agent Code*'] == imd, 'Biz Mix*'].str.upper()
                if not lobs.apply(lambda x: any(t in x for t in toks)).any():
                    sus.append(imd)
            if sus:
                cnt = int(imds.isin(sus).sum())
                F.add(L, 'S06b', 'Product-specific volume column used on an IMD that has no such product', WARN, cnt,
                      column=p, actual=f"IMDs: {', '.join(sus[:6])}; Biz Mix: {', '.join(sorted(set(s.loc[on & s['Parent Agent Code*'].isin(sus), 'Biz Mix*'])))}",
                      fix='Probably the wrong target column chosen manually for this Volume Consideration.')
    # STD-GRID columns must follow the STD-GRID rules for Agent Group + Biz Mix
    SR = std_rules(rules)
    vc_targets = set(user_map(rules.get('volume_consideration_map')).values())
    std_targets = {x['target'] for x in SR if x['target'] in setm.columns} - vc_targets
    if std_targets and 'Agent Group Code*' in s.columns:
        one = setm.sum(axis=1) == 1
        pair_of = setm.idxmax(axis=1).where(one, '')
        combo = s['Agent Group Code*'] + '\x1f' + s['Biz Mix*']
        cache = {}
        for k in pd.unique(combo):
            ag, bm = k.split('\x1f', 1); t, i = resolve_std(SR, ag, bm)
            cache[k] = (t or '(no rule)', describe_rule(SR[i]) if i is not None else f'{ag} + {bm}: no rule')
        want = combo.map(lambda k: cache[k][0])
        bad = pair_of.isin(std_targets) & (pair_of != want)
        if bad.any():
            d = pd.DataFrame({'ag': s['Agent Group Code*'][bad], 'act': pair_of[bad], 'exp': want[bad], 'rule': combo[bad].map(lambda k: cache[k][1]), 'bm': s['Biz Mix*'][bad]})
            for (ag, a_, e_, rule), grp in d.groupby(['ag', 'act', 'exp', 'rule'], sort=False):
                F.add(L, 'S06c', 'Volume is in a STD-GRID column that does not match the rule for this Agent Group + Biz Mix', ERROR, len(grp),
                      column=f'{ag}: {", ".join(sorted(set(grp["bm"]))[:6])}', expected=e_, actual=a_,
                      examples=_lines(grp.index) + f'  | rule: {rule}',
                      fix='STD-GRID row in the wrong column → move LL/UL to the rule column. If the row is not STD-GRID, its Volume Consideration fell back to Total Gwp → map it in Step 4.')
    # percent bands
    pb = rules['percent_band']
    for p in rules['percent_pairs']:
        if p + ' Ll*' not in N.columns: continue
        ll, ul = N[p + ' Ll*'], N[p + ' Ul*']
        rng = (ll < pb['min']) | (ul > pb['max']) | (ll > pb['max'])
        if rng.any():
            F.add(L, 'S08', 'Percent band out of range', ERROR, int(rng.sum()), column=p, examples=_lines(s.index[rng]),
                  fix=f"Expected {pb['min']}..{pb['max']} (defaults 0 / 100.001). Check Column Defaults.")
        frac = (ul > 0) & (ul <= 1.001)
        if frac.any():
            F.add(L, 'S08b', 'Percent band looks un-scaled (fraction, not ×100)', ERROR, int(frac.sum()), column=p + ' Ul*',
                  actual=', '.join(sorted(set(s.loc[frac, p + ' Ul*']))[:5]), examples=_lines(s.index[frac]),
                  fix='Manual step: ×100 and UL +0.001 (e.g. 0.4 → 40.001).')
        no001 = ul.notna() & ((ul * 1000).round().mod(10) != 1) & ~frac
        if no001.any():
            F.add(L, 'S08c', "Percent UL does not end in .001", WARN, int(no001.sum()), column=p + ' Ul*',
                  actual=', '.join(sorted(set(s.loc[no001, p + ' Ul*']))[:5]), examples=_lines(s.index[no001]))
    # span
    rc = rules['rate_cell']
    if 'Span Outgo*' in s.columns and 'Span Prct*' in s.columns:
        out, pr = s['Span Outgo*'], N['Span Prct*']
        rate_outgos = set(rc.get('outgo_by_header_label', {}).values()) | {rc['gwp_outgo']}
        bad = ~out.isin(rate_outgos | {rc['irda_outgo']})
        if bad.any():
            F.add(L, 'S09', 'Unexpected Span Outgo value', ERROR, int(bad.sum()), actual=', '.join(sorted(set(out[bad]))[:6]), examples=_lines(s.index[bad]))
        irda = out == rc['irda_outgo']
        b1 = irda & ((pr - rc['irda_prct']).abs() > 1e-9)
        if b1.any():
            F.add(L, 'S09b', f"IRDA rows must have Span Prct = {rc['irda_prct']}", ERROR, int(b1.sum()),
                  actual=', '.join(sorted(set(s.loc[b1, 'Span Prct*']))[:5]), examples=_lines(s.index[b1]),
                  fix='IRDA Prct in Step 2 must give -0.1 AFTER the ×100 transform (i.e. -0.001), not -0.1.')
        gwp = out.isin(rate_outgos)
        b2 = gwp & ((pr <= 0) | (pr > rc['max_gwp_prct']))
        if b2.any():
            F.add(L, 'S09c', 'Rate out of range (0, 100]', ERROR, int(b2.sum()), actual=', '.join(sorted(set(s.loc[b2, 'Span Prct*']))[:5]), examples=_lines(s.index[b2]))
        b3 = gwp & (pr > 0) & (pr < rc['min_gwp_prct'])
        if b3.any():
            F.add(L, 'S09d', 'Rate below 1 — probably not multiplied by 100', ERROR, int(b3.sum()),
                  actual=', '.join(sorted(set(s.loc[b3, 'Span Prct*']))[:5]), examples=_lines(s.index[b3]),
                  fix='Column Transform "Span Prct* × 100" missing in Step 6.')
    # version id
    if 'Version Id*' in s.columns:
        vids = s['Version Id*'].value_counts()
        if len(vids) != 1:
            F.add(L, 'S10', 'Version Id is not a single value', ERROR, len(vids), actual='; '.join(f'{k} ×{v}' for k, v in vids.items()))
        if version_id and (len(vids) != 1 or vids.index[0] != version_id):
            F.add(L, 'S10b', 'Version Id differs from the expected one', ERROR, int((s['Version Id*'] != version_id).sum()),
                  expected=version_id, actual=', '.join(vids.index[:3]))
    # allowed values & patterns
    for c, allowed in rules.get('allowed_values', {}).items():
        if c in s.columns:
            bad = ~s[c].isin(allowed)
            if bad.any():
                F.add(L, 'S11', 'Value not in allowed list', ERROR, int(bad.sum()), column=c, expected=', '.join(allowed),
                      actual=', '.join(sorted(set(s[c][bad]))[:6]), examples=_lines(s.index[bad]))
    pat = rules.get('agent_code_pattern')
    if pat:
        for c in ('Parent Agent Code*',):
            bad = ~s[c].str.match(pat)
            if bad.any():
                F.add(L, 'S11b', 'Agent code has unexpected format', WARN, int(bad.sum()), column=c,
                      actual=', '.join(sorted(set(s[c][bad]))[:6]), examples=_lines(s.index[bad]))
    # product signature: (Biz Mix + attributes) must equal a known grid column
    attr = [c for c in rules['product_attribute_columns'] if c in s.columns]
    sigs = set()
    for m in rules['column_map']:
        d = {c: rules['defaults'].get(c, '') for c in attr}
        d.update({k: v for k, v in m['fields'].items() if k in attr})
        sigs.add((m['biz_mix'],) + tuple(fmt_num(to_num(d[c])) if to_num(d[c]) is not None else d[c] for c in attr))
    known_bm = {x[0] for x in sigs}
    keyed = KK[['Biz Mix*'] + attr].copy(); keyed['Biz Mix*'] = s['Biz Mix*']
    tup = list(map(tuple, keyed.to_numpy()))
    unknown_bm = keyed['Biz Mix*'][~keyed['Biz Mix*'].isin(known_bm)]
    if len(unknown_bm):
        F.add(L, 'S12', 'Biz Mix not defined in rules (signature check skipped for these)', WARN, len(unknown_bm),
              actual=', '.join(sorted(set(unknown_bm))[:10]),
              fix='If this is a different grid type (e.g. STP/TP) use its own rules file.')
    badsig = [i for i, t in enumerate(tup) if t[0] in known_bm and t not in sigs]
    if badsig:
        cnt = Counter(tup[i] for i in badsig)
        txt = '; '.join(f"{t[0]}: " + ', '.join(f'{a}={v}' for a, v in zip(attr, t[1:]) if v != rules['defaults'].get(a, '') and v != fmt_num(to_num(rules['defaults'].get(a))) ) + f' ×{k}' for t, k in cnt.most_common(5))
        F.add(L, 'S12b', 'Biz Mix + attribute combination does not match any grid column (wrong Fuel/NCB/CC/Bus/ToB/Age)', ERROR, len(badsig),
              actual=txt, examples=_lines(pd.Index(badsig)), fix='Check Step 3 extra fields for these columns.')
    # RTO
    if 'Rto Code*' in s.columns:
        g = s.groupby('Rto Cluster*')['Rto Code*'].nunique()
        if (g > 1).any():
            F.add(L, 'S13', 'Same Rto Cluster has different Rto Code lists', ERROR, int((g > 1).sum()), actual=', '.join(g.index[g > 1][:8]))
        anyc = (s['Rto Code*'] == 'ANY') & (s['Rto Cluster*'] != 'ANY')
        if anyc.any():
            F.add(L, 'S13b', "Rto Code is 'ANY' while a cluster is given", ERROR, int(anyc.sum()), examples=_lines(s.index[anyc]),
                  fix='RTO file/sheet not picked up in Step 5.')
        if rto_master is not None:
            clus = s[['Rto Cluster*', 'Rto Code*']].drop_duplicates()
            allcodes = ','.join(sorted({x for v in rto_master.values() for x in v.split(',')}))
            for _, r in clus.iterrows():
                k = r['Rto Cluster*'].upper()
                if k == 'ANY': continue
                if k not in rto_master:
                    cnt = int((s['Rto Cluster*'] == r['Rto Cluster*']).sum())
                    F.add(L, 'S13c', 'Cluster not in RTO master', ERROR, cnt, column=r['Rto Cluster*'],
                          actual='portal fell back to ALL RTO codes' if r['Rto Code*'] == allcodes else r['Rto Code*'][:60])
                elif rto_master[k] != r['Rto Code*']:
                    want, got = set(rto_master[k].split(',')), set(r['Rto Code*'].split(','))
                    F.add(L, 'S13d', 'Rto Code list differs from RTO master', ERROR, int((s['Rto Cluster*'] == r['Rto Cluster*']).sum()),
                          column=r['Rto Cluster*'], expected=f'missing: {sorted(want-got)[:6]}', actual=f'extra: {sorted(got-want)[:6]}')
        br = rules.get('blocked_rtos')
        if br and br.get('codes'):
            u = s[['Rto Cluster*', 'Rto Code*']].drop_duplicates()
            hits = []
            for _, r in u.iterrows():
                codes = set(r['Rto Code*'].split(','))
                for b in br['codes']:
                    if b in codes: hits.append(f"{b} in {r['Rto Cluster*']}")
            if hits:
                F.add(L, 'S13e', 'Blocked RTOs (Generic Guidelines) still present in Rto Code', br.get('severity', WARN), len(hits),
                      actual=', '.join(hits), fix='Confirm whether blocked RTOs must be removed from the grid or are blocked elsewhere.')
    # duplicates
    dup = s.duplicated(keep=False)
    if dup.any():
        F.add(L, 'S14', 'Exact duplicate rows (same rate twice — remove one; some uploaders reject duplicates)', WARN,
              int(s.duplicated().sum()), examples=_lines(s.index[dup]))
    keycols = [c for c in s.columns if c not in ('Span Prct*', 'Span Outgo*')]
    kd = s.duplicated(subset=keycols, keep=False) & ~dup
    if kd.any():
        F.add(L, 'S14b', 'Conflicting rows: same criteria, different Span', ERROR, int(kd.sum()), examples=_lines(s.index[kd]),
              fix='Two source cells map to the same criteria — usually two grid columns given identical extra fields.')


def _lines(idx, k=6):
    idx = list(idx)[:k]
    return 'csv lines ' + ', '.join(str(int(i) + 2) for i in idx) if idx else ''


# ════════════════════════════════════════════════════════════════════════════
# LAYER 3 — PORTAL CONFIG LINT
# ════════════════════════════════════════════════════════════════════════════

def lint_config(cfg, L, sheet, rules, F):
    C = 'CONFIG'
    if L is not None:
        if cfg.get('sheet_name') != sheet:
            F.add(C, 'C01', 'Config sheet_name differs from the sheet being checked', WARN, 1, expected=sheet, actual=cfg.get('sheet_name'),
                  fix='Fine if you re-used the config, but confirm Step 2 points at the right sheet.')
        hr = [int(cfg.get(k, 0)) for k in ('header_row1', 'header_row2', 'header_row3')]
        want = [L.h1 + 1, L.h2 + 1, L.h3 + 1]
        if hr != want:
            F.add(C, 'C02', 'Header rows do not match the sheet layout', ERROR, 1, expected=want, actual=hr)
        if int(cfg.get('data_start_row', 0)) != L.data_start + 1:
            F.add(C, 'C02b', 'Data start row does not match the sheet', ERROR, 1, expected=L.data_start + 1, actual=cfg.get('data_start_row'))
        first_rate = min((r['col'] for r in L.rate_cols), default=None)
        if first_rate is not None and int(cfg.get('rate_cols_start', 0)) != first_rate + 1:
            F.add(C, 'C02c', 'Rate columns start column does not match the sheet', ERROR, 1, expected=first_rate + 1, actual=cfg.get('rate_cols_start'))
        meta_cfg = {'imd_code': 'mc_imd_code', 'imd_name': 'mc_imd_name', 'rel_code': 'mc_rel', 'imd_type': 'mc_imd_type',
                    'vol_remark': 'mc_vol_rem', 'uw_cluster': 'mc_cluster'}
        for k, ck in meta_cfg.items():
            if L.meta.get(k) is not None and int(cfg.get(ck, 0) or 0) != L.meta[k] + 1:
                F.add(C, 'C03', f'Meta column position wrong: {k}', ERROR, 1, expected=L.meta[k] + 1, actual=cfg.get(ck))
        if L.pairs.get('vol'):
            for ck, want in (('mc_vol_ll', L.pairs['vol'][0] + 1), ('mc_vol_ul', L.pairs['vol'][1] + 1)):
                if int(cfg.get(ck, 0) or 0) != want:
                    F.add(C, 'C03', f'Meta column position wrong: {ck}', ERROR, 1, expected=want, actual=cfg.get(ck))
        # column config vs actual headers
        hdr = {r['col'] + 1: r for r in L.rate_cols}
        cmap = {(nkey(m['parent']), nkey(m['sub'])): m for m in rules['column_map']}
        enabled = {c['col_idx']: c for c in cfg.get('colCfg', []) if c.get('enabled')}
        for ci, h in hdr.items():
            c = enabled.get(ci)
            if not c:
                F.add(C, 'C04', 'Rate column in sheet is not enabled in Step 3', ERROR, 1, column=f"{get_column_letter(ci)}: {h['parent']} | {h['sub']}")
                continue
            if (nkey(c.get('parent_biz')), nkey(c.get('sub_label'))) != (nkey(h['parent']), nkey(h['sub'])):
                F.add(C, 'C04b', 'Step 3 column points at a different header (columns shifted?)', ERROR, 1, column=get_column_letter(ci),
                      expected=f"{h['parent']} | {h['sub']}", actual=f"{c.get('parent_biz')} | {c.get('sub_label')}")
            if h['outgo'] != cfg.get('norm_outgo'):
                F.add(C, 'C04d', f"Column is on {h['label']} basis but portal writes Span Outgo = {cfg.get('norm_outgo')} for every column", ERROR, 1,
                      column=f"{get_column_letter(ci)}: {h['parent']} | {h['sub']}", expected=h['outgo'], actual=cfg.get('norm_outgo'),
                      fix='Portal has one Normal Outgo value per run: set Span Outgo manually for these rows (or extend the portal with a per-column outgo).')
            m = cmap.get((nkey(h['parent']), nkey(h['sub'])))
            if m:
                if c.get('biz_mix_output') != m['biz_mix']:
                    F.add(C, 'C05', 'Biz Mix output differs from rules', ERROR, 1, column=f"{h['parent']} | {h['sub']}", expected=m['biz_mix'], actual=c.get('biz_mix_output'))
                if (c.get('extra_fields') or {}) != m['fields']:
                    F.add(C, 'C05b', 'Extra fields differ from rules', ERROR, 1, column=f"{h['parent']} | {h['sub']}", expected=m['fields'], actual=c.get('extra_fields'))
        for ci in enabled:
            if ci not in hdr:
                F.add(C, 'C04c', 'Enabled Step 3 column has no GWP header in the sheet', ERROR, 1, column=f'{get_column_letter(ci)} ({enabled[ci].get("display")})')
        # extra meta columns
        groups = {k: v for k, v in L.pairs.items() if v}
        seen = Counter(int(e['col_idx']) for e in cfg.get('extraMetaCols', []))
        for col, k in seen.items():
            if k > 1:
                F.add(C, 'C06', 'Same source column used for two extra-meta outputs', ERROR, k, column=get_column_letter(col),
                      actual=', '.join(e['label'] for e in cfg['extraMetaCols'] if int(e['col_idx']) == col))
        for e in cfg.get('extraMetaCols', []):
            ci = int(e['col_idx']) - 1
            if L.meta.get('bizmix_label') == ci:
                F.add(C, 'C06b', 'Text label column (Biz Mix consideration) mapped as a numeric value', ERROR, 1, column=e['label'],
                      fix='The label decides WHICH Prct Vol pair gets the Biz Mix Lower/Upper; it is not a value. Needs the manual routing step.')
            if groups.get('bizmix') and ci in groups['bizmix'] and not e['label'].startswith(tuple(rules['bizmix_consideration_map'].values())):
                pass
            if groups.get('bizmix') and ci in groups['bizmix']:
                F.add(C, 'C06c', 'Biz Mix % band is hard-wired to one Prct Vol pair', WARN, 1, column=e['label'],
                      actual=f"source col {get_column_letter(ci+1)}",
                      fix='Rows whose Biz Mix consideration is TRACTOR OLD / PCV 3W NEW will land in the wrong pair — manual re-routing required.')
            if ' Ll*' in e['label'] and groups and any(ci == p[1] for p in groups.values()):
                F.add(C, 'C06d', 'LL output fed from an "Upper" column', ERROR, 1, column=e['label'], actual=get_column_letter(ci + 1))
            if ' Ul*' in e['label'] and groups and any(ci == p[0] for p in groups.values()):
                F.add(C, 'C06d', 'UL output fed from a "Lower" column', ERROR, 1, column=e['label'], actual=get_column_letter(ci + 1))
        # vol considerations / imd types in sheet vs config
        vrem_c, it_c = L.meta.get('vol_remark'), L.meta.get('imd_type')
        vals = Counter(); its = Counter()
        for r in L.rows[L.data_start:]:
            if vrem_c is not None and vrem_c < len(r) and r[vrem_c] not in (None, '', '-'): vals[str(r[vrem_c]).strip()] += 1
            if it_c is not None and it_c < len(r) and r[it_c]: its[str(r[it_c]).strip()] += 1
        vg = {nkey(v['vol_rem']) for v in cfg.get('vgRows', [])}
        for v, k in vals.items():
            if nkey(v) != 'STD-GRID' and nkey(v) not in vg:
                F.add(C, 'C07', 'Volume Consideration in sheet not mapped in Step 4 (falls to Total Gwp)', ERROR, k, actual=v,
                      expected=next((t for kk, t in rules['volume_consideration_map'].items() if not kk.startswith('_') and nkey(kk) == nkey(v)), 'no rule either'))
        ag = {nkey(a['imd_type']) for a in cfg.get('agentRows', [])}
        ig = {nkey(a['imd_type']) for a in cfg.get('imdGwpRows', [])}
        for v, k in its.items():
            if nkey(v) not in ag:
                F.add(C, 'C08', 'IMD Type not in Agent Group map', ERROR, k, actual=v)
            if nkey(v) not in ig and nkey(v) != 'IMD TYPE':
                F.add(C, 'C08b', 'IMD Type not in IMD→GWP map (std-grid uses keyword fallback)', WARN, k, actual=v)
    # vgRows / imdGwpRows target sanity
    out_cols = [f['col'] for f in cfg.get('outputFormat', [])]
    for src, rowsk in (('vgRows', 'vol_rem'), ('imdGwpRows', 'imd_type')):
        for r in cfg.get(src, []):
            ll, ul = r.get('ll_col', ''), r.get('ul_col', '')
            if not ll.endswith(' Ll*') or not ul.endswith(' Ul*') or ll[:-4] != ul[:-4]:
                F.add(C, 'C09', f'{src}: LL/UL target columns are not a matching pair', ERROR, 1, column=r.get(rowsk),
                      actual=f'LL→{ll}  UL→{ul}', fix='UL target must be the "… Ul*" column of the same pair.')
            for t in (ll, ul):
                if t and out_cols and t not in out_cols:
                    F.add(C, 'C09b', f'{src}: target column not in Output Format (value will be dropped)', ERROR, 1, column=r.get(rowsk), actual=t)
    rim = {nkey(k): v for k, v in user_map(rules['imd_type_map']).items()}
    SR = std_rules(rules)
    portal_std = {}
    for r in cfg.get('imdGwpRows', []):
        ag = (rim.get(nkey(r['imd_type'])) or {}).get('agent_group')
        portal_std[nkey(r['imd_type'])] = r.get('ll_col', '')[:-4]
        dflt = next((x['target'] for x in SR if norm_ag(x.get('agent_group')) == norm_ag(ag) and x.get('biz_mix_match', 'any') in ('any', '', None)), None)
        if dflt and r.get('ll_col', '')[:-4] != dflt:
            F.add(C, 'C09c', 'IMD→GWP map (std-grid) differs from the default STD-GRID rule', ERROR, 1, column=r['imd_type'], expected=dflt, actual=r.get('ll_col'))
    if L is not None and SR:
        def portal_target(it):
            if nkey(it) in portal_std: return portal_std[nkey(it)]
            u = it.upper()
            if 'PRIME' in u: return cfg.get('prime_ll', '')[:-4]
            if 'KEY' in u and 'BROK' in u: return cfg.get('keybrok_ll', '')[:-4]
            return cfg.get('agency_ll', '')[:-4]
        cmap2 = {(nkey(m['parent']), nkey(m['sub'])): m['biz_mix'] for m in rules['column_map']}
        rcols = [(r['col'], cmap2.get((nkey(r['parent']), nkey(r['sub'])))) for r in L.rate_cols]
        vc, itc = L.meta.get('vol_remark'), L.meta.get('imd_type')
        manual = Counter()
        for row in L.rows[L.data_start:]:
            if vc is None or vc >= len(row) or nkey(row[vc]) != 'STD-GRID': continue
            it = str(row[itc] or '').strip() if itc is not None and itc < len(row) else ''
            ag = (rim.get(nkey(it)) or {}).get('agent_group', it)
            pt = portal_target(it)
            for c, bm in rcols:
                if bm is None or c >= len(row) or row[c] in (None, ''): continue
                t, i = resolve_std(SR, ag, bm)
                if t and t != pt:
                    manual[(it, bm, pt, t, describe_rule(SR[i]), SR[i].get('biz_mix_match', 'any') in ('any', '', None))] += 1
        for (it, bm, pt, t, rule, is_default), k in sorted(manual.items(), key=lambda x: -x[1]):
            if is_default: continue   # already reported as C09c
            F.add(C, 'C15', 'Manual step needed: portal cannot apply a Biz-Mix-specific STD-GRID volume rule', WARN, k,
                  column=f'{it}: {bm}', expected=t, actual=pt, examples=f'rule: {rule}',
                  fix=f'After processing, move {pt} Ll*/Ul* → {t} Ll*/Ul* for these rows (the checker will verify it).')
    for k in ('std_ll', 'prime_ll', 'agency_ll', 'keybrok_ll'):
        uk = k.replace('_ll', '_ul')
        if cfg.get(k, '')[:-4] != cfg.get(uk, '')[:-4]:
            F.add(C, 'C09d', 'Fallback GWP LL/UL pair mismatch', ERROR, 1, column=k, actual=f"{cfg.get(k)} / {cfg.get(uk)}")
    # outputs & defaults
    exp_cols = rules['expected_columns']
    miss = [c for c in exp_cols if c not in out_cols]
    if miss:
        F.add(C, 'C10', 'Output Format is missing template columns', ERROR, len(miss), actual=', '.join(miss))
    defs = {r['col']: r['val'] for r in cfg.get('colDefaults', [])}
    for c in exp_cols:
        if c in ('Span Prct*', 'Span Outgo*', 'Version Id*'): continue
        want = rules['defaults'].get(c)
        if c not in defs:
            F.add(C, 'C11', 'Column has no default (blank / "-" will leak through)', ERROR, 1, column=c, expected=want)
        elif want is not None and str(defs[c]) != str(want):
            F.add(C, 'C11b', 'Column default differs from rules', ERROR, 1, column=c, expected=want, actual=defs[c])
    # transforms & IRDA value
    rc = rules['rate_cell']
    tr = [t for t in cfg.get('colTransforms', []) if t.get('col') == 'Span Prct*']
    mult = 1.0
    for t in tr:
        if t.get('op') == 'multiply': mult *= float(t.get('value', 1))
    if abs(mult - rc['scale']) > 1e-9:
        F.add(C, 'C12', 'Span Prct ×100 transform missing/wrong', ERROR, 1, expected=f"multiply {rc['scale']}", actual=f'net ×{mult}')
    try:
        if abs(float(cfg.get('irda_prct')) * mult - rc['irda_prct']) > 1e-9:
            F.add(C, 'C12b', 'IRDA Prct value gives the wrong final number after transforms', ERROR, 1,
                  expected=rc['irda_prct'], actual=f"{cfg.get('irda_prct')} × {mult} = {float(cfg.get('irda_prct'))*mult}")
    except (TypeError, ValueError):
        F.add(C, 'C12b', 'IRDA Prct value is not numeric', ERROR, 1, actual=cfg.get('irda_prct'))
    if cfg.get('mode') != 'special':
        F.add(C, 'C13', 'Mode is not "special" (IRDA cells will be skipped instead of producing rows)', ERROR, 1, actual=cfg.get('mode'))
    for k, v in (('irda_outgo', rc['irda_outgo']), ('norm_outgo', rc['gwp_outgo'])):
        if cfg.get(k) != v:
            F.add(C, 'C13b', f'{k} differs from rules', ERROR, 1, expected=v, actual=cfg.get(k))
    ign = {nkey(v) for v in str(cfg.get('ignore_values', '')).split('\n') if v.strip()}
    want = {nkey(v) for v in rc['irda_trigger_values']}
    if ign != want:
        F.add(C, 'C13c', 'IRDA trigger list differs from rules', ERROR, 1, expected=sorted(want), actual=sorted(ign))
    vid = next((r['val'] for r in cfg.get('staticFields', []) if r.get('col') == 'Version Id*'), None)
    if L is not None and vid:
        title = ' '.join(str(v) for r in L.rows[:max(L.h1, 1)] for v in r if v)
        months = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
        tm = next((m for m in months if re.search(m, title, re.I)), None)
        vm = next((m for m in months if m in vid.lower()), None)
        if tm and vm and tm != vm:
            F.add(C, 'C14', 'Version Id month does not match the grid title month', WARN, 1, expected=f'…{tm}…', actual=vid,
                  fix='Update Version Id in Step 6 if this is a new version.')



# ════════════════════════════════════════════════════════════════════════════
# CHECK CATALOG  (id → layer, default severity, description).  Every check can be switched off
# or given another severity in the rules file:  "checks": {"S13e": {"enabled": false}, "S14": {"severity": "ERROR"}}
# ════════════════════════════════════════════════════════════════════════════
CHECK_CATALOG = [
 ('S01','OUTPUT','ERROR','Expected template columns missing'), ('S01b','OUTPUT','ERROR','Unexpected columns present'),
 ('S01c','OUTPUT','WARN','Column order differs from template'), ('S02','OUTPUT','ERROR','Blank cells'),
 ('S03','OUTPUT','ERROR',"Placeholder values left ('-', nan, #N/A)"), ('S04','OUTPUT','ERROR','Non-numeric value in a numeric column'),
 ('S04b','OUTPUT','WARN','Floating-point artefacts (too many decimals)'), ('S05','OUTPUT','ERROR','Lower limit >= Upper limit'),
 ('S06','OUTPUT','ERROR','More than one volume column populated in a row'),
 ('S06b','OUTPUT','WARN','Product-specific volume column on an IMD without that product'),
 ('S06c','OUTPUT','ERROR','Volume in the wrong STD-GRID column for Agent Group + Biz Mix'),
 ('S06i','OUTPUT','INFO','Volume column usage by Agent Group (summary)'),
 ('S07','OUTPUT','ERROR','Volume sentinel on the wrong side (LL=999999999 / UL=-999999999)'),
 ('S08','OUTPUT','ERROR','Percent band out of range'), ('S08b','OUTPUT','ERROR','Percent band un-scaled (fraction, not x100)'),
 ('S08c','OUTPUT','WARN','Percent UL does not end in .001'), ('S09','OUTPUT','ERROR','Unexpected Span Outgo value'),
 ('S09b','OUTPUT','ERROR','IRDA rows must carry the IRDA Span Prct'), ('S09c','OUTPUT','ERROR','Rate out of range (0, 100]'),
 ('S09d','OUTPUT','ERROR','Rate below 1 (x100 missing)'), ('S10','OUTPUT','ERROR','Version Id not a single value'),
 ('S10b','OUTPUT','ERROR','Version Id differs from expected'), ('S11','OUTPUT','ERROR','Value not in allowed list'),
 ('S11b','OUTPUT','WARN','Agent code format'), ('S12','OUTPUT','WARN','Biz Mix not defined in rules'),
 ('S12b','OUTPUT','ERROR','Biz Mix + attributes match no grid column'), ('S13','OUTPUT','ERROR','Same cluster with different Rto Code lists'),
 ('S13b','OUTPUT','ERROR',"Rto Code 'ANY' while cluster given"), ('S13c','OUTPUT','ERROR','Cluster not in RTO master'),
 ('S13d','OUTPUT','ERROR','Rto Code list differs from RTO master'), ('S13e','OUTPUT','WARN','Blocked RTOs still present'),
 ('S14','OUTPUT','WARN','Exact duplicate rows'), ('S14b','OUTPUT','ERROR','Conflicting rows (same criteria, different Span)'),
 ('R00','SOURCE','ERROR','Meta column not found in sheet'), ('R01','SOURCE','ERROR','Unrecognised rate cell value in source'),
 ('R02','SOURCE','ERROR','Rate column in sheet has no rule'), ('R02b','SOURCE','INFO','Rule column not present in sheet'),
 ('R03','SOURCE','ERROR','Volume Consideration has no rule'), ('R03b','SOURCE','ERROR','IMD Type has no Agent Group rule'),
 ('R03c','SOURCE','ERROR','Biz Mix consideration label has no rule'), ('R03d','SOURCE','ERROR','No STD-GRID rule for Agent Group + Biz Mix'),
 ('R04','SOURCE','ERROR','Row count differs from source'), ('R04i','SOURCE','INFO','Row count matches source'),
 ('R05','SOURCE','ERROR','Expected rows missing from output'), ('R06','SOURCE','ERROR','Output rows with no source cell'),
 ('R07','SOURCE','ERROR','Cell value differs from source/rules'), ('R08','SOURCE','WARN','Where unruled volumes were placed'),
 ('R09','SOURCE','WARN','RTO master sheet unreadable'), ('R10','SOURCE','WARN','Duplicated IMD block in source sheet'),
 ('C01','CONFIG','WARN','Config sheet name differs'), ('C02','CONFIG','ERROR','Header rows wrong'), ('C02b','CONFIG','ERROR','Data start row wrong'),
 ('C02c','CONFIG','ERROR','Rate column start wrong'), ('C03','CONFIG','ERROR','Meta column position wrong'),
 ('C04','CONFIG','ERROR','Rate column not enabled'), ('C04b','CONFIG','ERROR','Step-3 column points at another header'),
 ('C04c','CONFIG','ERROR','Enabled column has no rate header'), ('C04d','CONFIG','ERROR','Column outgo basis (OD) differs from portal outgo'),
 ('C05','CONFIG','ERROR','Biz Mix output differs from rules'), ('C05b','CONFIG','ERROR','Extra fields differ from rules'),
 ('C06','CONFIG','ERROR','Source column used twice in extra-meta'), ('C06b','CONFIG','ERROR','Text label column mapped as value'),
 ('C06c','CONFIG','WARN','Biz Mix % band hard-wired to one pair'), ('C06d','CONFIG','ERROR','LL/UL fed from Upper/Lower column'),
 ('C07','CONFIG','ERROR','Volume Consideration not mapped in Step 4'), ('C08','CONFIG','ERROR','IMD Type not in Agent Group map'),
 ('C08b','CONFIG','WARN','IMD Type not in IMD->GWP map'), ('C09','CONFIG','ERROR','LL/UL target columns not a pair'),
 ('C09b','CONFIG','ERROR','Target column not in Output Format'), ('C09c','CONFIG','ERROR','IMD->GWP map differs from default STD-GRID rule'),
 ('C09d','CONFIG','ERROR','Fallback GWP LL/UL mismatch'), ('C10','CONFIG','ERROR','Output Format missing columns'),
 ('C11','CONFIG','ERROR','Column has no default'), ('C11b','CONFIG','ERROR','Column default differs from rules'),
 ('C12','CONFIG','ERROR','Span Prct x100 transform missing'), ('C12b','CONFIG','ERROR','IRDA Prct gives wrong final value'),
 ('C13','CONFIG','ERROR','Mode is not special'), ('C13b','CONFIG','ERROR','Outgo values differ from rules'),
 ('C13c','CONFIG','ERROR','IRDA trigger list differs'), ('C14','CONFIG','WARN','Version Id month differs from grid month'),
 ('C15','CONFIG','WARN','Manual step needed for Biz-Mix-specific STD-GRID rule'),
]


def validate_rules(r):
    """Return (errors, warnings) for a rules dict. Errors block saving from the portal."""
    errs, warns = [], []
    for k in ('expected_columns', 'defaults', 'column_map', 'imd_type_map', 'volume_consideration_map',
              'bizmix_consideration_map', 'rate_cell', 'percent_band', 'volume_pairs', 'percent_pairs',
              'product_attribute_columns', 'sheet_layout'):
        if k not in r: errs.append(f'Missing section: {k}')
    if errs: return errs, warns
    cols = set(r['expected_columns'])
    pair_ok = lambda p: (p + ' Ll*') in cols and (p + ' Ul*') in cols
    for i, x in enumerate(r.get('std_grid_volume_rules', []), 1):
        if x.get('biz_mix_match', 'any') not in MATCH_TYPES: errs.append(f'STD-GRID rule {i}: unknown match type {x.get("biz_mix_match")}')
        if not pair_ok(x.get('target', '')): errs.append(f'STD-GRID rule {i}: target "{x.get("target")}" has no Ll*/Ul* pair in expected columns')
        if x.get('biz_mix_match') == 'regex':
            try: re.compile(x.get('biz_mix_value', ''))
            except re.error as e: errs.append(f'STD-GRID rule {i}: bad regex ({e})')
        if x.get('biz_mix_match', 'any') not in ('any', '', None) and not str(x.get('biz_mix_value', '')).strip():
            errs.append(f'STD-GRID rule {i}: match "{x.get("biz_mix_match")}" needs a Biz Mix value')
    seen_any = set()
    for i, x in enumerate(r.get('std_grid_volume_rules', []), 1):
        g = norm_ag(x.get('agent_group'))
        if g in seen_any and x.get('enabled', True):
            warns.append(f'STD-GRID rule {i} can never match: an earlier "any Biz Mix" rule for {x.get("agent_group")} wins first')
        if x.get('biz_mix_match', 'any') in ('any', '', None) and x.get('enabled', True): seen_any.add(g)
    for name, sec in (('Volume Consideration', 'volume_consideration_map'), ('Biz Mix consideration', 'bizmix_consideration_map')):
        for k, v in user_map(r[sec]).items():
            if not pair_ok(v): errs.append(f'{name} "{k}": target "{v}" has no Ll*/Ul* pair in expected columns')
    for p in r['volume_pairs'] + r['percent_pairs']:
        if not pair_ok(p): errs.append(f'Pair "{p}" has no Ll*/Ul* columns in expected columns')
    for m in r['column_map']:
        for k in m.get('fields', {}):
            if k not in cols: errs.append(f'Column map {m.get("parent")} | {m.get("sub")}: field "{k}" not an expected column')
        if not m.get('biz_mix'): errs.append(f'Column map {m.get("parent")} | {m.get("sub")}: Biz Mix empty')
    for k in r['defaults']:
        if k not in cols: warns.append(f'Default for "{k}" — not an expected column')
    for c in r['expected_columns']:
        if c not in r['defaults'] and c not in ('Span Prct*', 'Span Outgo*', 'Version Id*', 'Biz Mix*', 'Rto Code*'):
            warns.append(f'No default for expected column "{c}"')
    for k, v in user_map(r['imd_type_map']).items():
        if not isinstance(v, dict) or not v.get('agent_group'): errs.append(f'IMD Type "{k}": agent_group missing')
    ids = {c[0] for c in CHECK_CATALOG}
    for k, v in (r.get('checks') or {}).items():
        if k not in ids: warns.append(f'checks: unknown check id {k}')
        elif isinstance(v, dict) and v.get('severity') not in (None, '', ERROR, WARN, INFO): errs.append(f'checks.{k}: bad severity')
    if r.get('agent_code_pattern'):
        try: re.compile(r['agent_code_pattern'])
        except re.error as e: errs.append(f'agent_code_pattern: {e}')
    return errs, warns

# ════════════════════════════════════════════════════════════════════════════
# REPORT
# ════════════════════════════════════════════════════════════════════════════

def write_report(path, F, meta, mismatch_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    base = Font(name='Arial', size=10); bold = Font(name='Arial', size=10, bold=True); hfont = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    hfill = PatternFill('solid', start_color='1F3864')
    sev_fill = {ERROR: PatternFill('solid', start_color='F8CBAD'), WARN: PatternFill('solid', start_color='FFE699'), INFO: PatternFill('solid', start_color='DDEBF7')}
    ws = wb.active; ws.title = 'Summary'
    ws['A1'] = 'Grid Checker Report'; ws['A1'].font = Font(name='Arial', size=14, bold=True)
    r = 3
    for k, v in meta.items():
        ws.cell(r, 1, k).font = bold; ws.cell(r, 2, str(v)).font = base; r += 1
    errs = sum(1 for i in F.items if i['severity'] == ERROR); warns = sum(1 for i in F.items if i['severity'] == WARN)
    verdict = 'FAIL — do not upload' if errs else ('PASS with warnings' if warns else 'PASS')
    ws.cell(r, 1, 'Verdict').font = bold; c = ws.cell(r, 2, verdict); c.font = Font(name='Arial', size=12, bold=True, color='C00000' if errs else '375623'); r += 1
    ws.cell(r, 1, 'Findings').font = bold; ws.cell(r, 2, f'{errs} errors, {warns} warnings').font = base; r += 2
    hdr = ['Layer', 'Check ID', 'Check', 'Severity', 'Findings', 'Affected rows/items']
    for j, h in enumerate(hdr, 1):
        c = ws.cell(r, j, h); c.font = hfont; c.fill = hfill
    agg = defaultdict(lambda: [0, 0, None, None, None])
    for i in F.items:
        k = (i['layer'], i['check_id'], i['check'], i['severity'])
        agg[k][0] += 1; agg[k][1] += i['count']
    order = {ERROR: 0, WARN: 1, INFO: 2}
    for (layer, cid, chk, sev), (nf, cnt, *_ ) in sorted(agg.items(), key=lambda x: (order[x[0][3]], x[0][0], x[0][1])):
        r += 1
        for j, v in enumerate([layer, cid, chk, sev, nf, cnt], 1):
            c = ws.cell(r, j, v); c.font = base
        ws.cell(r, 4).fill = sev_fill[sev]
    for col, w in zip('ABCDEF', (16, 12, 70, 10, 10, 18)):
        ws.column_dimensions[col].width = w

    wd = wb.create_sheet('Details')
    cols = ['layer', 'check_id', 'severity', 'check', 'column', 'expected', 'actual', 'count', 'examples', 'fix']
    for j, h in enumerate(cols, 1):
        c = wd.cell(1, j, h.replace('_', ' ').title()); c.font = hfont; c.fill = hfill
    for i, it in enumerate(sorted(F.items, key=lambda x: (order[x['severity']], x['layer'], x['check_id'], -x['count'])), 2):
        for j, h in enumerate(cols, 1):
            c = wd.cell(i, j, it[h]); c.font = base; c.alignment = Alignment(wrap_text=h in ('examples', 'fix', 'actual', 'expected'), vertical='top')
        wd.cell(i, 3).fill = sev_fill[it['severity']]
    for col, w in zip('ABCDEFGHIJ', (10, 8, 9, 48, 28, 30, 40, 9, 55, 55)):
        wd.column_dimensions[col].width = w
    wd.freeze_panes = 'A2'; wd.auto_filter.ref = f'A1:J{len(F.items)+1}'

    if mismatch_rows:
        wm = wb.create_sheet('Mismatch samples')
        mc = ['source_cell', 'csv_line', 'column', 'expected', 'actual']
        for j, h in enumerate(mc, 1):
            c = wm.cell(1, j, h.replace('_', ' ').title()); c.font = hfont; c.fill = hfill
        for i, m in enumerate(mismatch_rows[:5000], 2):
            for j, h in enumerate(mc, 1):
                wm.cell(i, j, m[h]).font = base
        for col, w in zip('ABCDE', (40, 10, 32, 30, 30)):
            wm.column_dimensions[col].width = w
        wm.freeze_panes = 'A2'
    wb.save(path)


# ════════════════════════════════════════════════════════════════════════════

def run(csv, rules_path, source=None, sheet=None, config=None, version_id=None, report=None, quiet=False, csv_label=None):
    rules = rules_path if isinstance(rules_path, dict) else json.load(open(rules_path))
    rules_name = rules.get('_name') or (os.path.basename(rules_path) if isinstance(rules_path, str) else 'rules')
    F = Findings(rules.get('checks'))
    act = pd.read_csv(csv, dtype=str, keep_default_na=False)
    act.columns = [c.strip() for c in act.columns]
    L = rto_master = None; mism = []
    if source:
        rto_master = read_rto_master(source, rules)
        if rto_master is None:
            F.add('SOURCE', 'R09', 'RTO master sheet could not be read — RTO checks against master skipped', WARN, 1)
        if sheet:
            L = read_sheet(source, sheet, rules)
    check_output(act, rules, F, rto_master, version_id)
    if L is not None:
        exp, refs, vu = build_expected(L, sheet, rules, rto_master, F, version_id)
        mism = reconcile(exp, refs, vu, act, rules, F, version_id)
    if config:
        lint_config(config if isinstance(config, dict) else json.load(open(config)), L, sheet, rules, F)
    meta = {'Output CSV': csv_label or os.path.basename(csv), 'Rows': len(act), 'Rules': rules_name,
            'Source': f"{os.path.basename(source)} :: {sheet}" if source and sheet else '(not given — output checks only)',
            'Portal config': ('current portal state' if isinstance(config, dict) else os.path.basename(config)) if config else '(not given)'}
    F.mismatch = mism; F.meta = meta
    if report:
        write_report(report, F, meta, mism)
    if not quiet:
        e = [i for i in F.items if i['severity'] == ERROR]; w = [i for i in F.items if i['severity'] == WARN]
        print(f"\n{'='*78}\n{os.path.basename(csv)}  —  {len(act):,} rows  —  {len(e)} errors / {len(w)} warnings\n{'='*78}")
        for it in sorted(F.items, key=lambda x: ({ERROR: 0, WARN: 1, INFO: 2}[x['severity']], x['layer'], x['check_id'])):
            if it['severity'] == INFO: continue
            col = f" [{it['column']}]" if it['column'] else ''
            ea = ''
            if it['expected'] or it['actual']:
                ea = f"  exp={it['expected'][:50]!s} act={it['actual'][:70]!s}"
            print(f"{it['severity']:5} {it['layer']:6} {it['check_id']:5} {it['count']:>7,}  {it['check']}{col}{ea}")
    return F


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--csv', required=True); ap.add_argument('--rules', required=True)
    ap.add_argument('--source'); ap.add_argument('--sheet'); ap.add_argument('--config')
    ap.add_argument('--version-id'); ap.add_argument('--report')
    a = ap.parse_args()
    try:
        F = run(a.csv, a.rules, a.source, a.sheet, a.config, a.version_id, a.report)
    except SystemExit:
        raise
    except Exception as ex:
        import traceback; traceback.print_exc(); sys.exit(2)
    sys.exit(1 if F.has_errors() else 0)
